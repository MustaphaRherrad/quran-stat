"""Lecture des sources Office et exports locaux UTF-8."""

import csv
import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook


def read_paragraphs(path):
    return [paragraph.text for paragraph in Document(path).paragraphs]


def read_verses(path):
    """Lit le corpus public et vérifie ses indices globaux."""
    verses = []
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames != ["Verse_ID", "Verse_Text"]:
            raise ValueError("Colonnes attendues : Verse_ID, Verse_Text.")
        for index, row in enumerate(reader, 1):
            if row["Verse_ID"] != str(index) or not row["Verse_Text"]:
                raise ValueError(f"Verset absent ou indice incorrect à la ligne {index + 1}.")
            verses.append(row["Verse_Text"])
    if not verses:
        raise ValueError("Le corpus est vide.")
    return verses


def write_csv(path, header, rows):
    with Path(path).open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)


def write_xlsx(path, header, rows):
    book = Workbook(write_only=True)
    sheet = book.create_sheet("Résultats")
    sheet.append(header)
    for row in rows:
        sheet.append(list(row))
    book.save(path)


def write_json(path, value):
    Path(path).write_text(
        json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
